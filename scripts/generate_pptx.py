import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_pptx_path):
    print(f"Generating PPTX presentation at {output_pptx_path}...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)      # Deep Slate / Navy
    LIGHT_BG = RGBColor(251, 251, 253)  # Apple Off-white
    CARD_BG = RGBColor(255, 255, 255)
    ACCENT_BLUE = RGBColor(0, 102, 204)  # Apple Blue
    ACCENT_AMBER = RGBColor(245, 158, 11) # Gold/Amber
    TEXT_DARK = RGBColor(29, 29, 31)
    TEXT_MUTED = RGBColor(100, 116, 139)
    TEXT_WHITE = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 232, 240)

    slides_data = [
        # Slide 1: Title Slide (Dark Theme)
        {
            "theme": "dark",
            "title": "VDCET Official Website with AI-Powered RAG Knowledge Assistant",
            "subtitle": "Final-Year B.Tech Computer Engineering Project Presentation (2026–2027)",
            "bullets": [
                "Student Lead: Suraj Gajanan Pathade",
                "Team Members: Suraj Gajanan Pathade, Rohit Shirbhate, Khushal Hajare, Harshal Kalbhut, Sanket Nasare, Tejas Sawargaokar",
                "Project Guide: Rohini Ma’am | Department of Computer Engineering",
                "Institution: Vilasrao Deshmukh College of Engineering & Technology, Mouda (DTE Code: 04141)",
                "Affiliated University: Dr. Babasaheb Ambedkar Technological University (DBATU), Lonere",
            ],
        },
        # Slide 2: Introduction
        {
            "theme": "light",
            "title": "1. Introduction & Institutional Profile",
            "subtitle": "Vilasrao Deshmukh College of Engineering & Technology (VDCET), Mouda",
            "bullets": [
                "Approved by AICTE, New Delhi & Directorate of Technical Education (DTE), Maharashtra.",
                "Official DTE Institute Code: 04141 (Used for CAP Allotment Rounds).",
                "Offers B.Tech programs in Computer Engineering, AI & Data Science, Civil, and Mechanical Engineering.",
                "Project Goal: Build a minimal, Apple-style web portal + grounded RAG AI assistant + hidden CMS.",
                "Practical Impact: Designed as a production-ready digital gift to VDCET.",
            ],
        },
        # Slide 3: Problem Statement
        {
            "theme": "light",
            "title": "2. Problem Statement",
            "subtitle": "Challenges with Traditional Institutional Web Portals",
            "bullets": [
                "Visual & Information Clutter: Traditional Indian college sites display unorganized banners and redundant links.",
                "Poor Mobile Usability: Slow load times and broken layouts on mobile networks.",
                "No Conversational Assistance: Prospective CAP candidates must read multiple pages to locate specific facts.",
                "AI Hallucination Risk: Generic chatbots frequently generate fake fees, dates, or faculty stats.",
                "Complex Administration: Lack of a simple, secure CMS for non-technical staff to publish notices.",
            ],
        },
        # Slide 4: Existing vs. Proposed System
        {
            "theme": "light",
            "title": "3. Existing System vs. Proposed System",
            "subtitle": "Architectural Shift Towards Grounded Conversational AI",
            "bullets": [
                "Existing System: Static legacy pages, crowded navigation, manual phone/office query handling.",
                "Proposed System: Apple-inspired layout, mobile-first performance, zero public login clutter.",
                "Grounded RAG AI Assistant: Answers strictly from official VDCET knowledge base chunks.",
                "Source Attribution: Displays verified source badges beneath answers.",
                "Hidden CMS (/admin): Protected by HTTP-only JWT session cookies.",
            ],
        },
        # Slide 5: Objectives
        {
            "theme": "light",
            "title": "4. Project Objectives",
            "subtitle": "Target Technical & Operational Deliverables",
            "bullets": [
                "1. Develop a minimal, mobile-first Apple-style web platform for VDCET Mouda.",
                "2. Implement a RAG AI Knowledge Assistant grounded in verified VDCET facts.",
                "3. Ensure zero false data generation via strict fallback messages when knowledge is missing.",
                "4. Build a hidden administrative CMS at secret route /admin.",
                "5. Donate the complete, fully-tested project package to VDCET.",
            ],
        },
        # Slide 6: System Architecture
        {
            "theme": "dark",
            "title": "5. System Architecture",
            "subtitle": "End-to-End Modular Layer Breakdown",
            "bullets": [
                "Client Layer: Mobile & Desktop Web Browsers (Public Visitors & Admin).",
                "Frontend Framework: Next.js 14 App Router + TypeScript + Tailwind CSS.",
                "API Layer: Next.js API Routes (/api/public/notices, /api/ai-chat, /api/admin/*).",
                "Intelligence Engine: RAG Vector Scorer + Groq Llama-3-8B-Instant LLM.",
                "Database Layer: Modular Data Access Repository (vdcet-db.json / SQLite).",
            ],
        },
        # Slide 7: Technology Stack
        {
            "theme": "light",
            "title": "6. Technology Stack Overview",
            "subtitle": "Modern Open-Source Technical Foundation",
            "bullets": [
                "Frontend & Server-side: Next.js 14, React 18, TypeScript",
                "Styling & UI: Tailwind CSS + Glassmorphism Utilities + Lucide React Icons",
                "Database Repository: Modular File Database (vdcet-db.json) / SQLite",
                "AI RAG Engine: Groq SDK (groq-sdk) Llama-3 + TF-IDF Keyword Matcher",
                "Security & Auth: JSON Web Tokens (jsonwebtoken) + HTTP-only Cookies + Bcrypt",
            ],
        },
        # Slide 8: RAG Retrieval Pipeline
        {
            "theme": "light",
            "title": "7. RAG Knowledge Retrieval Pipeline",
            "subtitle": "Deterministic Grounding & Anti-Hallucination Logic",
            "bullets": [
                "Step 1: Visitor submits question on /ai-assistant interface.",
                "Step 2: Engine calculates TF-IDF relevance scores against knowledge base chunks.",
                "Step 3: If no chunk matches (score = 0), returns official fallback notification.",
                "Step 4: Top matching chunks injected into grounded LLM system prompt.",
                "Step 5: Generated answer displayed alongside official source attribution badges.",
            ],
        },
        # Slide 9: Public Website Layout
        {
            "theme": "light",
            "title": "8. Public Website Structure",
            "subtitle": "8 Core Homepage Sections & Focused Navigation",
            "bullets": [
                "Public Routes: /, /about, /academics, /admissions, /campus, /notices, /contact, /ai-assistant",
                "Homepage Sequence: Hero -> Programs -> About -> Latest Notices -> Campus -> AI -> Contact -> Footer",
                "Hero Banner: DTE Code 04141 badge, title, subtitle, and campus photography.",
                "Footer Credit: Subtly displays 'Made by Developer Suraj' at bottom.",
                "No Public Login: Zero user registration or ERP overhead.",
            ],
        },
        # Slide 10: Academic Programs
        {
            "theme": "light",
            "title": "9. Academic B.Tech Programs (DTE: 04141)",
            "subtitle": "Four Specialized 4-Year Engineering Branches",
            "bullets": [
                "Computer Engineering (CSE): Software development, web architecture, algorithms.",
                "Artificial Intelligence & Data Science (AI & DS): Machine learning, analytics, neural networks.",
                "Civil Engineering (CIVIL): Structural design, surveying, hydraulics, material testing.",
                "Mechanical Engineering (MECH): Thermodynamics, CAD/CAM modeling, industrial manufacturing.",
                "All programs approved by DTE Maharashtra & AICTE, New Delhi.",
            ],
        },
        # Slide 11: Admissions & CAP Guide
        {
            "theme": "light",
            "title": "10. Admissions & CAP Process Guide",
            "subtitle": "Clear Guidelines for Maharashtra State Allotment Candidates",
            "bullets": [
                "DTE Allotment Code: Use Institute Code 04141 during CAP option form entry.",
                "Eligibility: 10+2 / HSC with PCM subjects + valid MHT-CET or JEE Main score.",
                "Step-by-Step CAP Guide: 1. CET Score -> 2. Option Form (04141) -> 3. Document Verification.",
                "Required Documents: HSC/SSC Marksheets, Allotment Letter, TC, Caste/Validity/Income certificates.",
                "Official Portal Link: Direct access to State CET Cell Maharashtra portal.",
            ],
        },
        # Slide 12: Database Design & ER Diagram
        {
            "theme": "dark",
            "title": "11. Database Design & ER Diagram",
            "subtitle": "Modular Entity Relationships",
            "bullets": [
                "Notices Entity: id, title, category, content, pdf_url, status, date",
                "Events Entity: id, title, date, location, description, image_url, status",
                "Documents Entity: id, title, category, file_url, file_size",
                "Knowledge Base Entity: id, title, content, source_name, category",
                "Website Content Entity: key, value, updated_at",
                "Status Support: Draft | Published | Archived controls.",
            ],
        },
        # Slide 13: Hidden Admin CMS
        {
            "theme": "light",
            "title": "12. Hidden Admin CMS (/admin)",
            "subtitle": "Isolated Administrative Management Portal",
            "bullets": [
                "Secret Route: /admin (Zero links or buttons on public website).",
                "Authentication: Password validation issuing HTTP-only JWT cookies.",
                "Management Tabs: Notices, Events, Documents, AI Knowledge Base, Website Content/Contacts.",
                "Status Controls: Draft, Published, Archived editorial controls.",
                "Strict Scope: No ERP, fees, payroll, or student portal complexity.",
            ],
        },
        # Slide 14: AI Assistant Live Demonstration Flow
        {
            "theme": "light",
            "title": "13. AI Assistant Query Demonstration",
            "subtitle": "Grounded Answers vs. Unfound Query Fallbacks",
            "bullets": [
                "Valid Query: 'What departments are available?'",
                "-> Answer: Lists CSE, AI&DS, Civil, Mech + Source Badge [Source: DTE Approved List].",
                "Unfound Query: 'Who is president of Mars?'",
                "-> Fallback Response: 'I couldn't find this information in the official VDCET knowledge base. Please contact the college office.'",
                "Grounding Guarantee: Zero hallucinated fees, dates, or rankings.",
            ],
        },
        # Slide 15: Empirical Testing Results
        {
            "theme": "light",
            "title": "14. Empirical System Test Matrix",
            "subtitle": "All 17 Systematic Test Cases Executed & Verified PASS",
            "bullets": [
                "TypeScript Typecheck: npx tsc --noEmit -> 0 Errors (PASS)",
                "Production Build: npm run build -> Compiled Successfully (PASS)",
                "Public Routes: GET /, /about, /academics, /admissions, /campus, /notices, /contact -> 200 OK (PASS)",
                "Admin Security: Access /admin -> Renders Login Card; Invalid Password -> 401 Unauthorized (PASS)",
                "RAG AI Query & Fallback: Valid query returns sources; Unfound query returns fallback (PASS)",
            ],
        },
        # Slide 16: Security & Privacy
        {
            "theme": "light",
            "title": "15. Security & Privacy Architecture",
            "subtitle": "Protecting System Data & Administrative Routes",
            "bullets": [
                "Hidden Route: /admin hidden from public navigation and search engines.",
                "Protected API Endpoints: All /api/admin/* routes verify JWT session cookies.",
                "Secret Management: API keys and secrets stored server-side in .env.local.",
                "XSS & Injection Safeguards: Strict React string escaping and sanitized inputs.",
            ],
        },
        # Slide 17: Future Scope
        {
            "theme": "light",
            "title": "16. Future Scope & Enhancements",
            "subtitle": "Potential Roadmap for Next Phases",
            "bullets": [
                "Multilingual AI Assistant: Regional language support (Marathi & Hindi queries).",
                "Voice Assistant: Integrated speech-to-text input for AI chat.",
                "Automated Document OCR: Direct PDF circular ingestion into the RAG knowledge base.",
                "Cloud Database Migration: Seamless transition to Supabase / PostgreSQL.",
            ],
        },
        # Slide 18: Conclusion
        {
            "theme": "dark",
            "title": "17. Conclusion & Acknowledgements",
            "subtitle": "Project Completion & Contribution to VDCET Mouda",
            "bullets": [
                "Successfully built an Apple-style web portal, RAG AI assistant, and hidden CMS.",
                "Tested with 100% verified results across all functional requirements.",
                "Ready for official deployment and donation to VDCET Mouda.",
                "Thank you to Project Guide Rohini Ma’am, HOD, and Computer Engineering Faculty!",
                "Questions & Viva Discussion",
            ],
        },
    ]

    for idx, slide_info in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)

        # Background color shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.line.fill.background()
        if slide_info["theme"] == "dark":
            bg.fill.solid()
            bg.fill.fore_color.rgb = DARK_BG
            title_color = TEXT_WHITE
            subtitle_color = ACCENT_AMBER
            body_color = RGBColor(226, 232, 240)
        else:
            bg.fill.solid()
            bg.fill.fore_color.rgb = LIGHT_BG
            title_color = TEXT_DARK
            subtitle_color = ACCENT_BLUE
            body_color = RGBColor(51, 65, 85)

        # Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.5))
        tf = title_box.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = slide_info["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(26 if idx > 0 else 30)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color

        if "subtitle" in slide_info and slide_info["subtitle"]:
            p_sub = tf.add_paragraph()
            p_sub.text = slide_info["subtitle"]
            p_sub.font.name = "Arial"
            p_sub.font.size = Pt(16)
            p_sub.font.color.rgb = subtitle_color
            p_sub.space_before = Pt(6)

        # Content Card / Bullet Box
        card_top = Inches(2.2)
        card_height = Inches(4.7)
        
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8),
            card_top,
            Inches(11.7),
            card_height,
        )
        if slide_info["theme"] == "dark":
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(30, 41, 59)
            card.line.color.rgb = RGBColor(51, 65, 85)
        else:
            card.fill.solid()
            card.fill.fore_color.rgb = CARD_BG
            card.line.color.rgb = BORDER_COLOR

        content_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.4), Inches(11.1), Inches(4.3))
        ctf = content_box.text_frame
        ctf.word_wrap = True

        for b_idx, bullet in enumerate(slide_info["bullets"]):
            p = ctf.paragraphs[0] if b_idx == 0 else ctf.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.name = "Arial"
            p.font.size = Pt(16)
            p.font.color.rgb = body_color
            p.space_after = Pt(12)

    prs.save(output_pptx_path)
    print(f"Successfully generated PPTX: {output_pptx_path}")

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    pres_dir = os.path.join(base_dir, "Presentation")

    create_presentation(os.path.join(pres_dir, "VDCET-Project-Presentation.pptx"))
